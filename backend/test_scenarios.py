import io
import os
import sys
import zipfile
import requests
from fpdf import FPDF
from docx import Document
from PIL import Image

BASE_URL = "http://127.0.0.1:8000/api/sourcemark"

session = requests.Session()
_login = session.post("http://127.0.0.1:8000/api/auth/login", json={"email": "demo@trustsphere.ai", "password": "demo123"})
if _login.status_code == 200:
    session.headers.update({"Authorization": "Bearer " + _login.json()["access_token"]})

def test_all():
    print("=" * 60)
    print("STARTING SOURCEMARK 2.0 INTEGRATION TEST SUITE")
    print("=" * 60)

    # -------------------------------------------------------------
    # Scenario 1 & 2: TXT Registration and Verification (Identical vs Modified)
    # -------------------------------------------------------------
    print("\n--- [Scenario 1 & 2] TXT File Comparison ---")
    nonce = os.urandom(4).hex()
    orig_txt = f"TrustSphere AI establishes a cryptographically verifiable security layer. Nonce: {nonce}\nAll operations are recorded to EvidenceLock.\nThis is a baseline configuration."
    mod_txt = f"TrustSphere AI establishes a cryptographically verifiable security layer. Nonce: {nonce}\nAll operations are recorded to EvidenceLock.\nThis is a MODIFIED configuration with enhanced integrity."

    txt_filename = f"policy_doc_{os.urandom(3).hex()}.txt"

    # Register original
    r = session.post(f"{BASE_URL}/upload", files={"file": (txt_filename, orig_txt.encode("utf-8"), "text/plain")})
    assert r.status_code == 200, f"TXT upload failed: {r.text}"
    reg_id = r.json()["id"]
    reg_hash = r.json()["original_hash"]
    print(f"Registered TXT '{txt_filename}' (ID: {reg_id}, Hash: {reg_hash[:16]}...)")

    # Verify identical (Scenario 1)
    r = session.post(f"{BASE_URL}/verify", files={"file": (txt_filename, orig_txt.encode("utf-8"), "text/plain")})
    assert r.status_code == 200, f"TXT identical verify failed: {r.text}"
    res = r.json()
    assert res["verification_status"] == "MATCH"
    assert res["comparison"]["change_percentage"] == 0.0
    assert res["comparison"]["similarity_percentage"] == 100.0
    print(f"Scenario 1 PASSED: Identical TXT -> Status: {res['verification_status']}, Change: {res['comparison']['change_percentage']}%, Similarity: {res['comparison']['similarity_percentage']}%")

    # Verify modified (Scenario 2)
    r = session.post(f"{BASE_URL}/verify", files={"file": (txt_filename, mod_txt.encode("utf-8"), "text/plain")})
    assert r.status_code == 200, f"TXT modified verify failed: {r.text}"
    res = r.json()
    assert res["verification_status"] == "MISMATCH"
    assert res["comparison"]["change_percentage"] > 0
    assert len(res["comparison"]["modified"]) > 0 or len(res["comparison"]["added"]) > 0
    comp_id_txt = res["comparison"]["comparison_id"]
    print(f"Scenario 2 PASSED: Modified TXT -> Status: {res['verification_status']}, Change: {res['comparison']['change_percentage']}%, Method: {res['comparison']['comparison_method']}")
    print(f"  Diff Details: {res['comparison']['modified'] or res['comparison']['added']}")

    # -------------------------------------------------------------
    # Scenario 3: PDF File Comparison
    # -------------------------------------------------------------
    print("\n--- [Scenario 3] PDF File Comparison ---")
    pdf_filename = f"audit_spec_{os.urandom(3).hex()}.pdf"

    def make_pdf(text_lines):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        for line in text_lines:
            pdf.cell(0, 10, line, ln=True)
        return bytes(pdf.output())

    orig_pdf_bytes = make_pdf(["Confidential Audit Document - Page 1", "Clause A: Strict adherence to hash integrity.", "Clause B: Unmodified parameter list."])
    mod_pdf_bytes = make_pdf(["Confidential Audit Document - Page 1", "Clause A: Strict adherence to hash integrity.", "Clause B: MODIFIED parameter list with elevated privileges."])

    # Register PDF
    r = session.post(f"{BASE_URL}/upload", files={"file": (pdf_filename, orig_pdf_bytes, "application/pdf")})
    assert r.status_code == 200, f"PDF upload failed: {r.text}"
    print(f"Registered PDF '{pdf_filename}'")

    # Verify modified PDF
    r = session.post(f"{BASE_URL}/verify", files={"file": (pdf_filename, mod_pdf_bytes, "application/pdf")})
    assert r.status_code == 200, f"PDF verify failed: {r.text}"
    res = r.json()
    assert res["verification_status"] == "MISMATCH"
    assert "affected_pages" in res["comparison"]
    print(f"Scenario 3 PASSED: Modified PDF -> Status: {res['verification_status']}, Affected: {res['comparison']['affected_pages']}, Change: {res['comparison']['change_percentage']}%")

    # -------------------------------------------------------------
    # Scenario 4: DOCX File Comparison
    # -------------------------------------------------------------
    print("\n--- [Scenario 4] DOCX File Comparison ---")
    docx_filename = f"handbook_{os.urandom(3).hex()}.docx"

    def make_docx(paragraphs):
        doc = Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    orig_docx_bytes = make_docx(["Section 1: Architecture Overview", "TrustSphere AI provides high-assurance verification."])
    mod_docx_bytes = make_docx(["Section 1: Architecture Overview", "TrustSphere AI provides high-assurance verification.", "Section 2: New Added Policy Section."])

    # Register DOCX
    r = session.post(f"{BASE_URL}/upload", files={"file": (docx_filename, orig_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200, f"DOCX upload failed: {r.text}"
    print(f"Registered DOCX '{docx_filename}'")

    # Verify modified DOCX
    r = session.post(f"{BASE_URL}/verify", files={"file": (docx_filename, mod_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200, f"DOCX verify failed: {r.text}"
    res = r.json()
    assert res["verification_status"] == "MISMATCH"
    assert len(res["comparison"]["added"]) > 0 or len(res["comparison"]["modified"]) > 0
    print(f"Scenario 4 PASSED: DOCX Modified -> Status: {res['verification_status']}, Added: {res['comparison']['added']}, Change: {res['comparison']['change_percentage']}%")

    # -------------------------------------------------------------
    # Scenario 5: Image (PNG) Comparison
    # -------------------------------------------------------------
    print("\n--- [Scenario 5] Image (PNG) Comparison ---")
    img_filename = f"diagram_{os.urandom(3).hex()}.png"

    orig_img = Image.new("RGB", (100, 100), color=(255, 255, 255))
    orig_img_buf = io.BytesIO()
    orig_img.save(orig_img_buf, format="PNG")
    orig_img_bytes = orig_img_buf.getvalue()

    mod_img = orig_img.copy()
    # Draw a 20x20 red square in center
    for x in range(40, 60):
        for y in range(40, 60):
            mod_img.putpixel((x, y), (255, 0, 0))
    mod_img_buf = io.BytesIO()
    mod_img.save(mod_img_buf, format="PNG")
    mod_img_bytes = mod_img_buf.getvalue()

    # Register Image
    r = session.post(f"{BASE_URL}/upload", files={"file": (img_filename, orig_img_bytes, "image/png")})
    assert r.status_code == 200, f"Image upload failed: {r.text}"
    img_reg_id = r.json()["id"]
    print(f"Registered Image '{img_filename}' (ID: {img_reg_id})")

    # Verify modified Image
    r = session.post(f"{BASE_URL}/verify", files={"file": (img_filename, mod_img_bytes, "image/png")}, data={"original_file_id": img_reg_id})
    assert r.status_code == 200, f"Image verify failed: {r.text}"
    res = r.json()
    assert res["verification_status"] == "MISMATCH"
    bbox = res["comparison"]["structural_details"].get("bounding_box")
    print(f"Scenario 5 PASSED: Image Modified -> Status: {res['verification_status']}, Bounding Box: {bbox}, Changed Pixels: {res['comparison']['structural_details'].get('changed_pixels')}, Change: {res['comparison']['change_percentage']}%")
    assert "diff_preview_base64" in res["comparison"]["structural_details"]

    # -------------------------------------------------------------
    # Scenario 6: ZIP Archive Comparison
    # -------------------------------------------------------------
    print("\n--- [Scenario 6] ZIP Archive Comparison ---")
    zip_filename = f"bundle_{os.urandom(3).hex()}.zip"

    def make_zip(file_dict):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for fname, content in file_dict.items():
                z.writestr(fname, content)
        return buf.getvalue()

    orig_zip_bytes = make_zip({"readme.txt": b"Package v1.0", "data.csv": b"1,2,3"})
    mod_zip_bytes = make_zip({"readme.txt": b"Package v1.0", "data.csv": b"1,2,3", "extra.key": b"secret123"})

    # Register ZIP
    r = session.post(f"{BASE_URL}/upload", files={"file": (zip_filename, orig_zip_bytes, "application/zip")})
    assert r.status_code == 200, f"ZIP upload failed: {r.text}"
    zip_reg_id = r.json()["id"]
    print(f"Registered ZIP '{zip_filename}' (ID: {zip_reg_id})")

    # Verify modified ZIP
    r = session.post(f"{BASE_URL}/verify", files={"file": (zip_filename, mod_zip_bytes, "application/zip")}, data={"original_file_id": zip_reg_id})
    assert r.status_code == 200, f"ZIP verify failed: {r.text}"
    res = r.json()
    assert res["verification_status"] == "MISMATCH"
    assert "extra.key" in str(res["comparison"]["added"])
    print(f"Scenario 6 PASSED: ZIP Modified -> Added entries: {res['comparison']['added']}, Change: {res['comparison']['change_percentage']}%")

    # -------------------------------------------------------------
    # Scenario 7: Generic Binary Comparison
    # -------------------------------------------------------------
    print("\n--- [Scenario 7] Generic Binary Comparison ---")
    bin_filename = f"firmware_{os.urandom(3).hex()}.dat"
    orig_bin = bytes([(i + int(os.urandom(1)[0])) % 256 for i in range(1000)])
    mod_bin = bytes([(b + 5) % 256 if 200 <= idx < 300 else b for idx, b in enumerate(orig_bin)])

    # Register Binary
    r = session.post(f"{BASE_URL}/upload", files={"file": (bin_filename, orig_bin, "application/octet-stream")})
    assert r.status_code == 200, f"Binary upload failed: {r.text}"
    bin_reg_id = r.json()["id"]
    print(f"Registered Binary '{bin_filename}' (ID: {bin_reg_id})")

    # Verify modified Binary
    r = session.post(f"{BASE_URL}/verify", files={"file": (bin_filename, mod_bin, "application/octet-stream")}, data={"original_file_id": bin_reg_id})
    assert r.status_code == 200, f"Binary verify failed: {r.text}"
    res = r.json()
    assert res["verification_status"] == "MISMATCH"
    assert res["comparison"]["structural_details"]["heuristic"] is True
    print(f"Scenario 7 PASSED: Binary Heuristic -> Byte Divergence: {res['comparison']['structural_details'].get('byte_divergence_percentage')}%, Change: {res['comparison']['change_percentage']}%")

    # -------------------------------------------------------------
    # Scenario 8: Corrupted / Malformed File
    # -------------------------------------------------------------
    print("\n--- [Scenario 8] Corrupted / Malformed File Handling ---")
    corrupt_filename = f"corrupt_{os.urandom(3).hex()}.pdf"
    corrupt_bytes = b"%PDF-1.4\nTRUNCATED_AND_INVALID_CONTENT_GARBAGE\x00\xff"

    # Register original valid PDF
    r = session.post(f"{BASE_URL}/upload", files={"file": (corrupt_filename, orig_pdf_bytes, "application/pdf")})
    assert r.status_code == 200
    corrupt_reg_id = r.json()["id"]

    # Verify with corrupt bytes - should fall back gracefully to binary, not crash with 500
    r = session.post(f"{BASE_URL}/verify", files={"file": (corrupt_filename, corrupt_bytes, "application/pdf")}, data={"original_file_id": corrupt_reg_id})
    assert r.status_code == 200, f"Corrupt file crashed backend with status {r.status_code}: {r.text}"
    res = r.json()
    assert res["verification_status"] == "MISMATCH"
    print(f"Scenario 8 PASSED: Corrupted PDF handled gracefully -> Method: {res['comparison']['comparison_method']}, Change: {res['comparison']['change_percentage']}%")

    # -------------------------------------------------------------
    # Scenario 9: Misleading Extension
    # -------------------------------------------------------------
    print("\n--- [Scenario 9] Misleading Extension Handling ---")
    fake_ext_filename = f"image_as_pdf_{os.urandom(3).hex()}.pdf"
    # An image named .pdf
    r = session.post(f"{BASE_URL}/upload", files={"file": (fake_ext_filename, orig_img_bytes, "application/pdf")})
    assert r.status_code == 200
    fake_reg_id = r.json()["id"]

    r = session.post(f"{BASE_URL}/verify", files={"file": (fake_ext_filename, mod_img_bytes, "application/pdf")}, data={"original_file_id": fake_reg_id})
    assert r.status_code == 200, f"Misleading extension crashed: {r.text}"
    res = r.json()
    assert res["verification_status"] == "MISMATCH"
    print(f"Scenario 9 PASSED: Misleading extension handled gracefully -> Status: {res['verification_status']}")

    # -------------------------------------------------------------
    # Scenario 10: PDF Integrity & Change Audit Report Download
    # -------------------------------------------------------------
    print("\n--- [Scenario 10] PDF Report Generation & Download ---")
    comp_id = res["comparison"]["comparison_id"]
    r = session.get(f"{BASE_URL}/report/{comp_id}/download")
    assert r.status_code == 200, f"Report download failed: {r.text}"
    assert r.headers.get("content-type") == "application/pdf"
    assert r.content.startswith(b"%PDF"), "Downloaded file is not a valid PDF!"
    print(f"Scenario 10 PASSED: Downloaded {len(r.content):,} bytes valid PDF report for comparison '{comp_id}'.")

    print("\n" + "=" * 60)
    print("ALL 10 SOURCEMARK 2.0 TEST SCENARIOS PASSED WITH 100% SUCCESS!")
    print("=" * 60)

if __name__ == "__main__":
    test_all()
